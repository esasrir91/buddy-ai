"""ReportTools - generate professional deliverables (Markdown / HTML / DOCX / PDF).

Reports *are* the product in accounting, valuation and engineering, so this
toolkit turns structured content (a title + ordered sections) into a saved
document. Markdown and HTML are always available (pure Python). DOCX needs
``python-docx`` and PDF needs ``weasyprint``; both degrade gracefully.

Every generated report can carry a mandatory disclaimer footer so regulated
outputs always state that a qualified professional must review and sign off.
"""

from __future__ import annotations

import html
import json
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional

from buddy.tools import Toolkit


def _normalize_sections(sections: List[Dict[str, str]]) -> List[Dict[str, str]]:
    norm = []
    for s in sections or []:
        if isinstance(s, dict):
            norm.append({"heading": str(s.get("heading", "")), "body": str(s.get("body", ""))})
    return norm


class ReportTools(Toolkit):
    def __init__(self, base_path: Optional[str] = None, default_disclaimer: str = "", **kwargs):
        self.base_path = Path(base_path).resolve() if base_path else (Path.cwd() / "reports").resolve()
        self.default_disclaimer = default_disclaimer
        super().__init__(
            name="report_tools",
            tools=[self.generate_markdown, self.generate_html, self.generate_docx, self.generate_pdf],
            **kwargs,
        )

    def _safe(self, filename: str) -> Path:
        p = (self.base_path / filename).resolve()
        if self.base_path not in p.parents and p != self.base_path:
            raise ValueError("Path escapes the allowed base directory.")
        p.parent.mkdir(parents=True, exist_ok=True)
        return p

    def _md(self, title: str, sections: List[Dict[str, str]], subtitle: str, disclaimer: str) -> str:
        lines = [f"# {title}", ""]
        if subtitle:
            lines += [f"*{subtitle}*", ""]
        lines += [f"_Generated {datetime.now():%Y-%m-%d %H:%M}_", ""]
        for s in _normalize_sections(sections):
            if s["heading"]:
                lines += [f"## {s['heading']}", ""]
            lines += [s["body"], ""]
        d = disclaimer or self.default_disclaimer
        if d:
            lines += ["---", "", f"> **Disclaimer.** {d}", ""]
        return "\n".join(lines)

    def generate_markdown(self, filename: str, title: str, sections: List[Dict[str, str]], subtitle: str = "", disclaimer: str = "") -> str:
        """Write a Markdown report.

        Args:
            filename (str): Output .md path relative to the base directory.
            title (str): Report title.
            sections (List[Dict[str,str]]): Ordered sections, each {"heading": ..., "body": ...}.
            subtitle (str): Optional subtitle line.
            disclaimer (str): Optional disclaimer footer (falls back to the toolkit default).

        Returns:
            str: JSON with the saved path.
        """
        try:
            out = self._safe(filename if filename.endswith(".md") else filename + ".md")
            out.write_text(self._md(title, sections, subtitle, disclaimer), encoding="utf-8")
            return json.dumps({"saved": str(out), "format": "markdown"})
        except Exception as e:
            return json.dumps({"error": str(e)})

    def _html(self, title: str, sections: List[Dict[str, str]], subtitle: str, disclaimer: str) -> str:
        parts = [
            "<!doctype html><meta charset='utf-8'>",
            "<style>body{font-family:Georgia,serif;max-width:820px;margin:40px auto;color:#1a1a1a;line-height:1.55}"
            "h1{font-size:26px}h2{font-size:18px;border-bottom:1px solid #ddd;padding-bottom:4px;margin-top:28px}"
            ".sub{color:#666;font-style:italic}.meta{color:#999;font-size:12px}"
            ".disc{margin-top:32px;padding:12px 16px;background:#f6f6f6;border-left:3px solid #c00;font-size:13px;color:#444}</style>",
            f"<h1>{html.escape(title)}</h1>",
        ]
        if subtitle:
            parts.append(f"<p class='sub'>{html.escape(subtitle)}</p>")
        parts.append(f"<p class='meta'>Generated {datetime.now():%Y-%m-%d %H:%M}</p>")
        for s in _normalize_sections(sections):
            if s["heading"]:
                parts.append(f"<h2>{html.escape(s['heading'])}</h2>")
            body = html.escape(s["body"]).replace("\n", "<br>")
            parts.append(f"<div>{body}</div>")
        d = disclaimer or self.default_disclaimer
        if d:
            parts.append(f"<div class='disc'><b>Disclaimer.</b> {html.escape(d)}</div>")
        return "\n".join(parts)

    def generate_html(self, filename: str, title: str, sections: List[Dict[str, str]], subtitle: str = "", disclaimer: str = "") -> str:
        """Write a styled HTML report (also used as the source for PDF).

        Args:
            filename (str): Output .html path relative to the base directory.
            title (str): Report title.
            sections (List[Dict[str,str]]): Ordered {"heading","body"} sections.
            subtitle (str): Optional subtitle.
            disclaimer (str): Optional disclaimer footer.

        Returns:
            str: JSON with the saved path.
        """
        try:
            out = self._safe(filename if filename.endswith(".html") else filename + ".html")
            out.write_text(self._html(title, sections, subtitle, disclaimer), encoding="utf-8")
            return json.dumps({"saved": str(out), "format": "html"})
        except Exception as e:
            return json.dumps({"error": str(e)})

    def generate_docx(self, filename: str, title: str, sections: List[Dict[str, str]], subtitle: str = "", disclaimer: str = "") -> str:
        """Write a Microsoft Word (.docx) report (requires python-docx).

        Args:
            filename (str): Output .docx path relative to the base directory.
            title (str): Report title.
            sections (List[Dict[str,str]]): Ordered {"heading","body"} sections.
            subtitle (str): Optional subtitle.
            disclaimer (str): Optional disclaimer footer.

        Returns:
            str: JSON with the saved path, or a hint to install python-docx.
        """
        try:
            from docx import Document  # type: ignore
        except Exception:
            return json.dumps({"error": "python-docx not installed. Run: pip install python-docx (or use generate_markdown/html)."})
        try:
            doc = Document()
            doc.add_heading(title, level=0)
            if subtitle:
                doc.add_paragraph(subtitle).italic = True
            doc.add_paragraph(f"Generated {datetime.now():%Y-%m-%d %H:%M}")
            for s in _normalize_sections(sections):
                if s["heading"]:
                    doc.add_heading(s["heading"], level=1)
                doc.add_paragraph(s["body"])
            d = disclaimer or self.default_disclaimer
            if d:
                doc.add_paragraph()
                p = doc.add_paragraph()
                run = p.add_run("Disclaimer. " + d)
                run.italic = True
            out = self._safe(filename if filename.endswith(".docx") else filename + ".docx")
            doc.save(str(out))
            return json.dumps({"saved": str(out), "format": "docx"})
        except Exception as e:
            return json.dumps({"error": str(e)})

    def generate_pdf(self, filename: str, title: str, sections: List[Dict[str, str]], subtitle: str = "", disclaimer: str = "") -> str:
        """Write a PDF report (requires weasyprint; falls back to HTML otherwise).

        Args:
            filename (str): Output .pdf path relative to the base directory.
            title (str): Report title.
            sections (List[Dict[str,str]]): Ordered {"heading","body"} sections.
            subtitle (str): Optional subtitle.
            disclaimer (str): Optional disclaimer footer.

        Returns:
            str: JSON with the saved path, or a fallback HTML path if weasyprint is missing.
        """
        markup = self._html(title, sections, subtitle, disclaimer)
        try:
            from weasyprint import HTML  # type: ignore
        except Exception:
            try:
                out = self._safe(filename.replace(".pdf", "") + ".html")
                out.write_text(markup, encoding="utf-8")
                return json.dumps({"saved": str(out), "format": "html", "note": "weasyprint not installed; wrote HTML instead. Install weasyprint for PDF."})
            except Exception as e:
                return json.dumps({"error": str(e)})
        try:
            out = self._safe(filename if filename.endswith(".pdf") else filename + ".pdf")
            HTML(string=markup).write_pdf(str(out))
            return json.dumps({"saved": str(out), "format": "pdf"})
        except Exception as e:
            return json.dumps({"error": str(e)})


__all__ = ["ReportTools"]
